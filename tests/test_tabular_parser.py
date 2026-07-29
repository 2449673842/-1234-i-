import json
import os
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime

from openpyxl import Workbook

try:
    import xlwt
except ImportError:
    xlwt = None


PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PARSER = os.path.join(PROJECT_ROOT, "renderer", "tabular_parser.py")


class TabularParserTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.xlsx_path = os.path.join(self.temp_dir.name, "sample.xlsx")
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Primary"
        sheet.append(["样本名称", "浓度（µg/L）", "采样日期"])
        sheet.append(["alpha", 1.5, datetime(2026, 7, 10, 12, 30, 0)])
        sheet.append(["beta", None, None])
        secondary = workbook.create_sheet("Secondary")
        secondary.append(["ignored"])
        secondary.append(["not-first-sheet"])
        workbook.save(self.xlsx_path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def run_parser(self, mode, input_path=None, limit=None, extra_args=None):
        command = [
            sys.executable,
            PARSER,
            "--input",
            input_path or self.xlsx_path,
            "--mode",
            mode,
        ]
        if limit is not None:
            command.extend(["--limit", str(limit)])
        if extra_args:
            command.extend(extra_args)
        completed = subprocess.run(
            command,
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=30,
        )
        return json.loads(completed.stdout)

    def test_metadata_uses_first_sheet(self):
        result = self.run_parser("metadata")
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["columns"], ["样本名称", "浓度（µg/L）", "采样日期"])
        self.assertEqual(result["row_count"], 2)
        self.assertNotIn("rows", result)

    def test_records_normalize_nulls_and_dates(self):
        result = self.run_parser("records")
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["rows"][0]["样本名称"], "alpha")
        self.assertEqual(result["rows"][0]["浓度（µg/L）"], 1.5)
        self.assertTrue(result["rows"][0]["采样日期"].startswith("2026-07-10T12:30:00"))
        self.assertIsNone(result["rows"][1]["浓度（µg/L）"])
        self.assertIsNone(result["rows"][1]["采样日期"])

    def test_records_limit_does_not_change_total_row_count(self):
        result = self.run_parser("records", limit=1)
        self.assertEqual(result["status"], "success")
        self.assertEqual(result["row_count"], 2)
        self.assertEqual(len(result["rows"]), 1)
        self.assertEqual(result["rows"][0]["样本名称"], "alpha")

    @unittest.skipUnless(xlwt is not None, "xlwt is required to generate an .xls fixture")
    def test_xls_matches_xlsx_metadata_and_record_semantics(self):
        xls_path = os.path.join(self.temp_dir.name, "unicode-sample.xls")
        workbook = xlwt.Workbook()
        sheet = workbook.add_sheet("Primary")
        headers = ["样本名称", "浓度（µg/L）", "采样日期"]
        for column, value in enumerate(headers):
            sheet.write(0, column, value)
        date_style = xlwt.XFStyle()
        date_style.num_format_str = "yyyy-mm-dd hh:mm:ss"
        sheet.write(1, 0, "alpha")
        sheet.write(1, 1, 1.5)
        sheet.write(1, 2, datetime(2026, 7, 10, 12, 30, 0), date_style)
        sheet.write(2, 0, "beta")
        workbook.save(xls_path)

        metadata = self.run_parser("metadata", input_path=xls_path)
        records = self.run_parser("records", input_path=xls_path)
        self.assertEqual(metadata["status"], "success")
        self.assertEqual(metadata["columns"], headers)
        self.assertEqual(metadata["row_count"], 2)
        self.assertEqual(records["status"], "success")
        self.assertEqual(records["rows"][0]["浓度（µg/L）"], 1.5)
        self.assertTrue(records["rows"][0]["采样日期"].startswith("2026-07-10T12:30:00"))
        self.assertIsNone(records["rows"][1]["浓度（µg/L）"])
        self.assertIsNone(records["rows"][1]["采样日期"])

    def test_corrupt_workbook_returns_structured_error(self):
        corrupt_path = os.path.join(self.temp_dir.name, "corrupt.xlsx")
        with open(corrupt_path, "wb") as handle:
            handle.write(b"not-an-xlsx")
        completed = subprocess.run(
            [sys.executable, PARSER, "--input", corrupt_path, "--mode", "metadata"],
            check=True,
            capture_output=True,
            text=True,
            encoding="utf-8",
            timeout=30,
        )
        result = json.loads(completed.stdout)
        self.assertEqual(result["status"], "error")
        self.assertIn("Workbook parsing failed", result["message"])

    def test_shape_budget_is_checked_before_materializing_records(self):
        result = self.run_parser("metadata", extra_args=["--max-columns", "2"])
        self.assertEqual(result["status"], "error")
        self.assertIn("more than 2 columns", result["message"])


if __name__ == "__main__":
    unittest.main()
