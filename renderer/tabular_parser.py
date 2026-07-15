import argparse
import json
import math
import os
import sys
import zipfile
from datetime import date, datetime

import pandas as pd
import openpyxl
import xlrd


def normalize_value(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except Exception:
        pass
    if isinstance(value, (datetime, date, pd.Timestamp)):
        return value.isoformat()
    if hasattr(value, "item"):
        try:
            value = value.item()
        except Exception:
            pass
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def validate_shape(row_count, column_count, max_rows, max_columns, max_cells, materialized_rows=None):
    if column_count > max_columns:
        raise ValueError(f"Workbook has more than {max_columns} columns")
    if row_count > max_rows:
        raise ValueError(f"Workbook has more than {max_rows} data rows")
    cells_row_count = row_count if materialized_rows is None else min(row_count, materialized_rows)
    if cells_row_count * column_count > max_cells:
        raise ValueError(f"Workbook has more than {max_cells} data cells")


def validate_xlsx_archive(file_path, max_entries, max_uncompressed_bytes, max_ratio):
    with zipfile.ZipFile(file_path, "r") as archive:
        entries = archive.infolist()
        if not entries or len(entries) > max_entries:
            raise ValueError("XLSX archive entry count exceeds the safety budget")
        names = {entry.filename.replace("\\", "/") for entry in entries}
        if "[Content_Types].xml" not in names or "xl/workbook.xml" not in names:
            raise ValueError("ZIP container is not an XLSX workbook")
        total_uncompressed = 0
        for entry in entries:
            name = entry.filename.replace("\\", "/")
            if name.startswith("/") or ".." in name.split("/"):
                raise ValueError("XLSX archive contains an unsafe internal path")
            total_uncompressed += int(entry.file_size)
            if total_uncompressed > max_uncompressed_bytes:
                raise ValueError("XLSX uncompressed size exceeds the safety budget")
            if entry.file_size > 0 and (
                entry.compress_size <= 0 or entry.file_size / entry.compress_size > max_ratio
            ):
                raise ValueError("XLSX compression ratio exceeds the safety budget")


def workbook_metadata(file_path, extension, limits, materialized_rows=None):
    engine = "xlrd" if extension == ".xls" else "openpyxl"
    if extension == ".xlsx":
        validate_xlsx_archive(
            file_path,
            limits["max_entries"],
            limits["max_uncompressed_bytes"],
            limits["max_compression_ratio"],
        )
        workbook = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        try:
            worksheet = workbook.worksheets[0]
            row_count = max(0, int(worksheet.max_row or 0) - 1)
            column_count = max(0, int(worksheet.max_column or 0))
        finally:
            workbook.close()
    else:
        workbook = xlrd.open_workbook(file_path, on_demand=True)
        try:
            worksheet = workbook.sheet_by_index(0)
            row_count = max(0, int(worksheet.nrows) - 1)
            column_count = max(0, int(worksheet.ncols))
        finally:
            workbook.release_resources()
    validate_shape(
        row_count,
        column_count,
        limits["max_rows"],
        limits["max_columns"],
        limits["max_cells"],
        materialized_rows,
    )
    header = pd.read_excel(file_path, sheet_name=0, engine=engine, nrows=0)
    columns = [str(column) for column in header.columns]
    for column in columns:
        if len(column) > limits["max_cell_chars"]:
            raise ValueError("Workbook header cell exceeds the safety budget")
    return columns, row_count


def parse_workbook(file_path, mode, limit=None, limits=None):
    extension = os.path.splitext(file_path)[1].lower()
    if extension not in {".xlsx", ".xls"}:
        raise ValueError("Only .xlsx and .xls files are supported")

    engine = "xlrd" if extension == ".xls" else "openpyxl"
    limits = limits or {
        "max_rows": 200_000,
        "max_columns": 512,
        "max_cells": 2_000_000,
        "max_cell_chars": 262_144,
        "max_entries": 2_048,
        "max_uncompressed_bytes": 256 * 1024 * 1024,
        "max_compression_ratio": 200,
    }
    columns, row_count = workbook_metadata(
        file_path,
        extension,
        limits,
        limit if mode == "records" and limit and limit > 0 else None,
    )
    result = {
        "status": "success",
        "columns": columns,
        "row_count": row_count,
        "sheet_index": 0,
    }
    if mode == "records":
        frame = pd.read_excel(
            file_path,
            sheet_name=0,
            engine=engine,
            nrows=limit if limit and limit > 0 else None,
        )
        rows = []
        for row in frame.to_dict(orient="records"):
            normalized = {}
            for column, value in row.items():
                normalized_value = normalize_value(value)
                if isinstance(normalized_value, str) and len(normalized_value) > limits["max_cell_chars"]:
                    raise ValueError("Workbook cell exceeds the safety budget")
                normalized[str(column)] = normalized_value
            rows.append(normalized)
        result["rows"] = rows
    return result


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", required=True)
    parser.add_argument("--mode", choices=["metadata", "records"], default="metadata")
    parser.add_argument("--limit", type=int, default=None)
    parser.add_argument("--max-rows", type=int, default=200_000)
    parser.add_argument("--max-columns", type=int, default=512)
    parser.add_argument("--max-cells", type=int, default=2_000_000)
    parser.add_argument("--max-cell-chars", type=int, default=262_144)
    parser.add_argument("--max-entries", type=int, default=2_048)
    parser.add_argument("--max-uncompressed-bytes", type=int, default=256 * 1024 * 1024)
    parser.add_argument("--max-compression-ratio", type=int, default=200)
    args = parser.parse_args()

    try:
        result = parse_workbook(args.input, args.mode, args.limit, {
            "max_rows": max(1, args.max_rows),
            "max_columns": max(1, args.max_columns),
            "max_cells": max(1, args.max_cells),
            "max_cell_chars": max(1, args.max_cell_chars),
            "max_entries": max(1, args.max_entries),
            "max_uncompressed_bytes": max(1, args.max_uncompressed_bytes),
            "max_compression_ratio": max(1, args.max_compression_ratio),
        })
    except Exception as error:
        result = {
            "status": "error",
            "message": f"Workbook parsing failed: {error}",
        }
    sys.stdout.write(json.dumps(result, ensure_ascii=False, separators=(",", ":")))


if __name__ == "__main__":
    main()
